"""Report whether Max Pain / GEX have predicted anything so far, using the Outcome sheet
that outcomes.py fills in. Read-only: it prints a report and changes nothing.

The key test is Max Pain vs the naive "price stays where it is" guess. If Max Pain doesn't
beat that, it isn't a signal. Small samples prove nothing either way (see MIN_N).
"""
import json
import os
import sys
from pathlib import Path
from statistics import mean

import openpyxl

from outcomes import GUESSED
from rules import RULES, stats
from update_gex import _mode
from update_workbooks import zone

HERE = Path(__file__).parent
CONFIG = Path(os.environ.get("MAXPAIN_CONFIG", HERE / "config.json"))
MIN_N = 30


def _core(records):
    n = len(records)
    return {
        "n": n,
        "toward_rate": sum(abs(r["settle"] - r["mp"]) < abs(r["start"] - r["mp"]) for r in records) / n,
        "mae_mp": mean(abs(r["settle"] - r["mp"]) for r in records),
        "mae_naive": mean(abs(r["settle"] - r["start"]) for r in records),
    }


def summarize(records, mid, edge):
    """records: dicts with start, mp, settle, net, range_pct - only rows that have a settle."""
    if not records:
        return {"n": 0}
    s = _core(records)
    s["range_by_gex"] = {}
    for label, test in (("NET GEX > 0", lambda v: v > 0), ("NET GEX < 0", lambda v: v < 0)):
        vals = [r["range_pct"] for r in records
                if isinstance(r["net"], (int, float)) and test(r["net"]) and r["range_pct"] is not None]
        s["range_by_gex"][label] = (len(vals), mean(vals) if vals else None)
    # which GEX reading tells volatility apart better: the single contract or every expiration
    s["range_by_mode"] = {}
    for src, key in (("single", "mode"), ("รวม", "mode_all")):
        for sign, label in (("+", "Positive"), ("-", "Negative")):
            vals = [r["range_pct"] for r in records
                    if _mode(r.get(key)) == sign and r["range_pct"] is not None]
            s["range_by_mode"][f"{label} ({src})"] = (len(vals), mean(vals) if vals else None)
    zones = {}
    for r in records:
        zones.setdefault(zone((r["mp"] - r["start"]) / r["start"], mid, edge), []).append(r)
    s["by_zone"] = {z: _core(rs) for z, rs in zones.items()}
    return s


def format_report(s):
    if not s["n"]:
        return "No outcomes with an expiry settle yet - nothing to analyze."
    lines = [f"Sample: {s['n']} expired contracts"]
    if s["n"] < MIN_N:
        lines.append(f"WARNING: fewer than {MIN_N} samples - these numbers are noise, not evidence.")
    better = "Max Pain" if s["mae_mp"] < s["mae_naive"] else "naive (no change)"
    lines += [
        "",
        f"Moved toward Max Pain by expiry: {s['toward_rate']:.0%}",
        f"Avg |settle - Max Pain|:  {s['mae_mp']:.1f}",
        f"Avg |settle - start|:     {s['mae_naive']:.1f}   (naive baseline)",
        f"Closer predictor so far:  {better}",
        "",
        "Realized range (high-low / start) by NET GEX sign:",
    ]
    for label, (n, avg) in s["range_by_gex"].items():
        lines.append(f"  {label}: " + (f"{avg:.2%} (n={n})" if avg is not None else "no data"))
    lines.append("Realized range by GEX mode, single contract vs all expirations (รวม):")
    for label, (n, avg) in s.get("range_by_mode", {}).items():
        lines.append(f"  {label}: " + (f"{avg:.2%} (n={n})" if avg is not None else "no data"))
    lines.append("")
    lines.append("By Max Pain zone at log time:")
    for z, c in s["by_zone"].items():
        lines.append(f"  {z}: n={c['n']}, toward {c['toward_rate']:.0%}, "
                     f"MP err {c['mae_mp']:.1f} vs naive {c['mae_naive']:.1f}")
    return "\n".join(lines)


def record(r):
    """One Outcome sheet row (values) -> the dict summarize()/rules.py read."""
    r = tuple(r) + (None,) * (29 - len(r))
    return {"start": r[3], "mp": r[4], "net": r[5], "settle": r[7], "range_pct": r[14],
            "hi": r[8], "lo": r[9], "entry": r[18], "change": r[19], "mode": r[20],
            "call_wall": r[22], "put_wall": r[23],
            "mode_all": r[25], "call_wall_all": r[27], "put_wall_all": r[28]}


def load_records(cfg):
    """(records, rows left out as guessed-month, mid, edge) from the Outcome sheet, or None."""
    wb = openpyxl.load_workbook(HERE / cfg["workbook"], data_only=True)
    if "Outcome" not in wb.sheetnames:
        return None
    calc = wb["Max Pain Calc"]
    mid = calc["F9"].value if isinstance(calc["F9"].value, (int, float)) else 0.02
    edge = calc["F10"].value if isinstance(calc["F10"].value, (int, float)) else 0.05
    records, guessed = [], 0
    for r in wb["Outcome"].iter_rows(min_row=2, values_only=True):
        if not (isinstance(r[7], (int, float)) and isinstance(r[3], (int, float)) and r[3]):
            continue
        if len(r) > 17 and str(r[17] or "").startswith(GUESSED):
            guessed += 1
            continue
        records.append(record(r))
    return records, guessed, mid, edge


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # Thai text on a cp1252 console
    loaded = load_records(json.load(open(CONFIG, encoding="utf-8")))
    if loaded is None:
        print("No Outcome sheet yet - run outcomes.py (it also runs as part of the daily update).")
        return
    records, guessed, mid, edge = loaded
    print(format_report(summarize(records, mid, edge)))
    print("\n" + format_paper(paper_trades(records)))
    if guessed:
        print(f"\n(Left out {guessed} row(s) whose futures month was guessed - see outcomes.GUESSED.)")


def paper_trades(records):
    """{rule name: stats} over the records that have an entry price (newer Outcome rows)."""
    usable = [r for r in records if r.get("entry") is not None]
    return {name: stats([t for t in map(rule, usable) if t]) for name, rule in RULES.items()}


def format_paper(by_rule):
    lines = ["Paper trades (R1/R2/R3 fixed 2026-10-03, R1-all/R2-all 2026-10-10; entry next-day open,"
             " exit expiry settle, stops worst-case):"]
    base = by_rule.get("Baseline: always long", {}).get("n", 0)
    if base < MIN_N:
        lines.append(f"WARNING: {base} usable contracts - fewer than {MIN_N}, these results are noise.")
    for name, s in by_rule.items():
        if not s["n"]:
            lines.append(f"  {name}: no trades yet")
            continue
        lines.append(f"  {name}: n={s['n']}, win {s['win_rate']:.0%}, total {s['total']:+.1f} pts, "
                     f"avg {s['avg']:+.1f}, worst {s['worst']:+.1f}, stopped {s['stopped']}")
    return "\n".join(lines)


if __name__ == "__main__":
    main()
