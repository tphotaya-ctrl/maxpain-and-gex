"""Backfill Log-sheet history for trade dates CME still has data for.

CME's own trade-date list is a short rolling window (currently ~5 trading days), so this
catches up a few missed days (machine off, expired login) - it is not a way to build long-run
history, since CME simply stops serving older days through these endpoints. Never touches
`OI Data` / `Gamma Data`, which only ever hold the most recent day's snapshot; only appends
to `Log`, and skips any date already logged (safe to run repeatedly).

    python backfill.py            # Max Pain workbook only
    python backfill.py --gex      # also backfill the GEX workbook (needs the QuikStrike login)

For each historical trade date T, the contract picked is the one the daily run would have
picked the morning after: `target` = T + 1 day (see target_for) - not whatever `config.json`'s
own `target` says, which usually means "today" and would be wrong for every day but the latest.
"""
import argparse
import json
import os
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from charts import rebuild_charts
from fetch_cme import cme_session, expiry_date, fetch
from fetch_gamma import fetch_gamma, qs_code
from update_gex import GEX_LOG, format_last_row, gamma_flip, gex_log, gex_status, levels, select_rows
from update_workbooks import LOG_FMT as MP_LOG_FMT
from update_workbooks import LOG_HEAD as MP_LOG_HEAD
from update_workbooks import max_pain
from util import save_atomic

HERE = Path(__file__).parent
CONFIG = Path(os.environ.get("MAXPAIN_CONFIG", HERE / "config.json"))


def _trade_dates(cfg):
    """All trade dates CME's own history currently covers, newest first."""
    with cme_session(cfg) as get:
        return [d["tradeDate"] for d in get("/CmeWS/mvc/Volume/TradeDates?exchange=CBOT&isProtected")]


def sort_by_date(ws):
    """Backfilled days land after newer ones; put data rows back in date order, so 'the last
    row' (alerts, daily summary) is always the newest day. Values only - cell formats stay
    per position, and every Log row uses the same formats."""
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    key = [(_day(r[0]), i) for i, r in enumerate(rows)]
    if key == sorted(key, key=lambda k: (str(k[0]), k[1])):
        return False
    order = [i for _, i in sorted(key, key=lambda k: (str(k[0]), k[1]))]
    for out_i, src_i in enumerate(order, start=2):
        for col, v in enumerate(rows[src_i], start=1):
            ws.cell(out_i, col).value = v
    return True


def target_for(td):
    """The `target` the daily run would have used for trade date td ('YYYYMMDD'): it runs
    the next morning, so it asks for the contract expiring the *next* calendar day (with the
    usual next-expiry fallback over weekends). Targeting td itself picks the contract that
    expired that day, which has zero OI at the close - so every weekday came back empty."""
    return str(datetime.strptime(td, "%Y%m%d").date() + timedelta(days=1))


def _day(v):
    return v.date() if isinstance(v, datetime) else v


def catch_up(cfg):
    """Daily-run hook: fill any trade date CME still has that's missing from either Log."""
    trade_dates = _trade_dates(cfg)
    backfill_max_pain(cfg, trade_dates)
    backfill_gex(cfg, trade_dates, retries=1)  # runs daily: a date QuikStrike lacks fails fast


def _existing_dates(path, sheet):
    if not Path(path).exists():
        return set()
    wb = openpyxl.load_workbook(path, read_only=True)
    try:  # read-only mode keeps the file open until close(); Windows then refuses the save
        if sheet not in wb.sheetnames:
            return set()
        return {r[0].date() for r in wb[sheet].iter_rows(min_row=2, values_only=True)
                if isinstance(r[0], datetime)}
    finally:
        wb.close()


def backfill_max_pain(cfg, trade_dates):
    path = HERE / cfg["workbook"]
    have = _existing_dates(path, "Log")
    wb = openpyxl.load_workbook(path)
    if "Log" not in wb.sheetnames:
        wb.create_sheet("Log").append(MP_LOG_HEAD)
    log = wb["Log"]
    added = 0
    for td in trade_dates:
        if datetime.strptime(td, "%Y%m%d").date() in have:
            continue
        target = target_for(td)
        try:
            d = fetch({**cfg, "trade_date": td, "target": target, "strike_min": 0, "strike_max": 10**9})
        except Exception as e:
            print(f"Max Pain: skip {td} ({type(e).__name__}: {e})")
            continue
        mp, low, valley = max_pain(d["strikes"])
        calls, puts = sum(s[1] for s in d["strikes"]), sum(s[2] for s in d["strikes"])
        total = d["call_total"] + d["put_total"]
        try:
            dte = (expiry_date(d["family"], d["label"]) - d["trade_date"]).days
        except (KeyError, ValueError, IndexError):  # monthly (AME) labels have no weekday rule
            dte = None
        log.append([d["trade_date"], f"{d['family']} {d['label']} ({d['code']})", dte, d["price"], mp,
                    (mp - d["price"]) / d["price"] if d["price"] else None, valley,
                    puts / calls if calls else None, (calls + puts) / total if total else None,
                    d["report"], datetime.now(), d["price_month"]])
        for cell, fmt in zip(log[log.max_row], MP_LOG_FMT):
            if fmt:
                cell.number_format = fmt
        added += 1
        print(f"Max Pain: logged {td} - {d['family']} {d['label']} maxpain={mp}")
    if added:
        sort_by_date(log)
        rebuild_charts(wb)
        save_atomic(wb, path)
    print(f"Max Pain backfill: {added} day(s) added, {len(trade_dates) - added} already had a row or failed")


def backfill_gex(cfg, trade_dates, retries=3):
    path = HERE / cfg["gex_workbook"]
    have = _existing_dates(path, GEX_LOG)
    wb = openpyxl.load_workbook(path)
    log = gex_log(wb)
    gc = wb["GEX Calc"] if "GEX Calc" in wb.sheetnames else None
    added = 0
    for td in trade_dates:
        if datetime.strptime(td, "%Y%m%d").date() in have:
            continue
        target = target_for(td)
        try:
            d = fetch({**cfg, "trade_date": td, "target": target, "strike_min": 0, "strike_max": 10**9})
            code = qs_code(d["family"], d["label"])
            if expiry_date(d["family"], d["label"]) < date.today():
                # QuikStrike only lists live contracts - once one expires its gamma is gone
                print(f"GEX: skip {td} ({code} has expired; QuikStrike no longer lists it)")
                continue
            g = fetch_gamma(cfg, code, d["trade_date"], retries=retries)
        except Exception as e:
            print(f"GEX: skip {td} ({type(e).__name__}: {e})")
            continue
        rows = select_rows(g)
        calls, puts = sum(r[1] for r in rows), sum(r[2] for r in rows)
        net = calls - puts
        log.append([d["trade_date"], code, d["price"], calls, puts, net, gex_status(rows),
                    gamma_flip(rows, d["price"]), max(rows, key=lambda r: r[1] + r[2])[0], datetime.now(),
                    *levels(rows, d["price"], gc)])
        format_last_row(log)
        added += 1
        print(f"GEX: logged {td} - {code} net={net:+.0f}")
    if added:
        sort_by_date(log)
        rebuild_charts(wb)
        save_atomic(wb, path)
    print(f"GEX backfill: {added} day(s) added, {len(trade_dates) - added} already had a row or failed")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--gex", action="store_true", help="also backfill the GEX workbook")
    args = ap.parse_args()
    cfg = json.load(open(CONFIG, encoding="utf-8"))
    trade_dates = _trade_dates(cfg)
    print(f"CME currently has {len(trade_dates)} trade date(s) available: {trade_dates[-1]}..{trade_dates[0]}")
    backfill_max_pain(cfg, trade_dates)
    if args.gex:
        backfill_gex(cfg, trade_dates)


if __name__ == "__main__":
    main()
