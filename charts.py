"""openpyxl drops every chart when it saves a workbook. rebuild_charts() puts back the
Max Pain line chart and the GEX V.4.1 bar chart, for whichever of their sheets the
workbook has, sized to the rows actually filled. Call it right before every save."""
from openpyxl.chart import BarChart, LineChart, Reference

MP_LAST = 201   # OI Data rows 2-201 feed Max Pain Calc
GEX_LAST = 152  # Gamma Data rows 2-152 feed GEX Calc (GEX V.4.1)


def filled(ws, last):
    return sum(1 for r in range(2, last + 1) if ws.cell(r, 1).value is not None)


def max_pain_chart(calc, n):
    calc._charts = []
    if not n:
        return
    ch = LineChart()
    ch.title, ch.legend = "PAIN รวม ตาม Strike", None
    ch.height, ch.width = 7.5, 26
    ch.add_data(Reference(calc, min_col=2, min_row=1, max_row=1 + n), titles_from_data=True)
    ch.set_categories(Reference(calc, min_col=1, min_row=2, max_row=1 + n))
    calc.add_chart(ch, "H2")


def gex_chart(gc, n):
    gc._charts = []
    if not n:
        return
    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title = "Gross Gamma ตาม Strike (แรง hedging)"
    ch.y_axis.title, ch.x_axis.title = "Gross Gamma", "Strike"
    ch.height, ch.width = 7.5, 15
    ch.add_data(Reference(gc, min_col=3, min_row=1, max_row=1 + n), titles_from_data=True)
    ch.set_categories(Reference(gc, min_col=1, min_row=2, max_row=1 + n))
    gc.add_chart(ch, "K1")


def rebuild_charts(wb):
    names = wb.sheetnames
    if "Max Pain Calc" in names and "OI Data" in names:
        max_pain_chart(wb["Max Pain Calc"], filled(wb["OI Data"], MP_LAST))
    if "GEX Calc" in names and "Gamma Data" in names:
        gex_chart(wb["GEX Calc"], filled(wb["Gamma Data"], GEX_LAST))
