"""One Telegram message after each daily run: today's Max Pain / GEX reading and which
paper rule (rules.py) is armed for the next session. Printed too, so run.log keeps it.

    python daily_summary.py     # rebuild and send from what the workbook holds now
"""
import json
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

from notify import send_telegram
from rules import MP_MIN_GAP
from update_gex import GEX_LOG, _mode
from update_workbooks import zone

HERE = Path(__file__).parent
CONFIG = Path(os.environ.get("MAXPAIN_CONFIG", HERE / "config.json"))


def _n(v, fmt="{:,.0f}"):
    return fmt.format(v) if isinstance(v, (int, float)) else "-"


def _day(v):
    return v.date() if isinstance(v, datetime) else v


def rule_lines(price, mp, mode, call_wall, put_wall, change):
    """What each paper rule would do next session, using the latest settle as the price
    (the real entry is the next open, so R1 is phrased as a condition on that open)."""
    m = _mode(mode)
    if m == "+" and call_wall is not None and put_wall is not None:
        mid = (call_wall + put_wall) / 2
        r1 = (f"R1: กรอบ {_n(put_wall)}-{_n(call_wall)} - เปิดต่ำกว่า {_n(mid)} ซื้อ (stop {_n(put_wall)}), "
              f"สูงกว่า ขาย (stop {_n(call_wall)}); นอกกรอบไม่เข้า")
    else:
        r1 = "R1: ไม่เข้า (ต้องเป็นโหมดบวกและมี Wall ครบ)"
    if m == "-" and change:
        if change > 0:
            stop = put_wall if put_wall is not None and price and put_wall < price else None
            r2 = f"R2: Long ตามราคาที่ขึ้น {change:+.1f}" + (f", stop {_n(stop)}" if stop else ", ไม่มี stop")
        else:
            stop = call_wall if call_wall is not None and price and call_wall > price else None
            r2 = f"R2: Short ตามราคาที่ลง {change:+.1f}" + (f", stop {_n(stop)}" if stop else ", ไม่มี stop")
    elif m == "-":
        r2 = "R2: ไม่เข้า (ไม่มีข้อมูลราคาเปลี่ยนแปลง)"
    else:
        r2 = "R2: ไม่เข้า (ต้องเป็นโหมดลบ)"
    if mp is not None and price and abs(mp - price) / price >= MP_MIN_GAP:
        r3 = f"R3: {'Long' if mp > price else 'Short'} ไปหา Max Pain {_n(mp)}"
    else:
        r3 = f"R3: ไม่เข้า (ห่าง Max Pain ไม่ถึง {MP_MIN_GAP:.1%})"
    return [r1, r2, r3]


def format_summary(mp_row, gex_row, mid, edge, change=None):
    """mp_row: last Max Pain Log row; gex_row: the GEX Log row for the same day, or None."""
    date_, contract, dte, price, mp, measure = mp_row[:6]
    report, month = mp_row[9], mp_row[11] if len(mp_row) > 11 else None
    z = zone(measure, mid, edge)
    lines = [
        f"MaxPain/GEX - ข้อมูล {_day(date_)} ({report})",
        f"สัญญา {contract} - DTE {dte}",
        f"ราคา {_n(price, '{:,.1f}')}" + (f" ({month})" if month else "")
        + (f" เปลี่ยน {change:+.1f}" if change else ""),
        f"Max Pain {_n(mp)}" + (f" ({measure:+.2%}, โซน{z})" if measure is not None else ""),
    ]
    mode = call_wall = put_wall = None
    if gex_row:
        g = tuple(gex_row) + (None,) * (14 - len(gex_row))
        mode, call_wall, put_wall = g[6], g[10], g[11]
        lines.append(f"GEX: {mode} - NET {_n(g[5], '{:+,.0f}')} - ความชัดเจน {_n(g[13], '{:.2f}')}")
        lines.append(f"Put Wall {_n(put_wall)} - Call Wall {_n(call_wall)} - Pin {_n(g[12])}")
    else:
        lines.append("GEX: ไม่มีข้อมูลวันนี้ (ข้ามไป - ดู run.log)")
    lines.append("")
    lines.append("กฎเทรดจำลอง (session ถัดไป):")
    lines += ["- " + s for s in rule_lines(price, mp, mode, call_wall, put_wall, change)]
    return "\n".join(lines)


def _last(ws):
    return [c.value for c in ws[ws.max_row]] if ws.max_row > 1 else None


def main(change=None):
    cfg = json.load(open(CONFIG, encoding="utf-8"))
    wb = openpyxl.load_workbook(HERE / cfg["workbook"], data_only=True)
    mp_row = _last(wb["Log"])
    if not mp_row:
        return None
    calc = wb["Max Pain Calc"]
    mid = calc["F9"].value if isinstance(calc["F9"].value, (int, float)) else 0.02
    edge = calc["F10"].value if isinstance(calc["F10"].value, (int, float)) else 0.05
    gwb = wb if cfg["gex_workbook"] == cfg["workbook"] else openpyxl.load_workbook(HERE / cfg["gex_workbook"])
    gex_row = _last(gwb[GEX_LOG]) if GEX_LOG in gwb.sheetnames else None
    if gex_row and _day(gex_row[0]) != _day(mp_row[0]):
        gex_row = None  # GEX was skipped today; don't report an older day's reading
    text = format_summary(mp_row, gex_row, mid, edge, change)
    print(text)
    send_telegram(text)
    return text


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
