"""One-time: build MaxPain_GEX.xlsx = the Max Pain workbook + GEX V.4.1's three sheets +
the old GEX workbook's Log (as 'GEX Log'). The source files are only opened read-only.

    python merge_workbooks.py            # refuses to overwrite an existing output
    python merge_workbooks.py --force

Sheets are copied by Excel itself (COM), which keeps charts, conditional formatting,
dropdowns and array formulas exactly; openpyxl can't copy sheets between workbooks. Then
openpyxl widens V.4.1's ranges (update_gex.extend_layout) and rebuilds both charts.
Afterwards point config.json's 'workbook' and 'gex_workbook' at the output.
"""
import argparse
import base64
import subprocess
import sys
from pathlib import Path

import openpyxl

from charts import rebuild_charts
from update_gex import GEX_LOG, extend_layout
from util import save_atomic

HERE = Path(__file__).parent
MAX_PAIN = HERE / "คำนวน_max_pain_10 (2).xlsx"
GEX_NEW = HERE / "GEX V.4.1.xlsx"
GEX_OLD = HERE / "คำนวน GEX 1.xlsx"  # only its Log history is carried over
OUT = HERE / "MaxPain_GEX.xlsx"
GEX_SHEETS = ["Gamma Data", "GEX Calc", "บันทึกโหมด"]


def _q(p):
    return str(p).replace("'", "''")


def excel_merge():
    names = ",".join(f"'{s}'" for s in GEX_SHEETS)
    ps = f"""
$ErrorActionPreference='Stop'
$m=[System.Reflection.Missing]::Value
$x=New-Object -ComObject Excel.Application
$x.Visible=$false; $x.DisplayAlerts=$false
try {{
  $out=$x.Workbooks.Open('{_q(MAX_PAIN)}',0,$true)
  $out.SaveAs('{_q(OUT)}',51)
  $g=$x.Workbooks.Open('{_q(GEX_NEW)}',0,$true)
  $g.Worksheets.Item([object[]]@({names})).Copy($m,$out.Worksheets.Item($out.Worksheets.Count))
  $old=$x.Workbooks.Open('{_q(GEX_OLD)}',0,$true)
  $old.Worksheets.Item('Log').Copy($m,$out.Worksheets.Item($out.Worksheets.Count))
  $out.Worksheets.Item($out.Worksheets.Count).Name='{GEX_LOG}'
  $links=$out.LinkSources(1)
  if ($links -is [Array]) {{
    for ($i=$links.GetLowerBound(0); $i -le $links.GetUpperBound(0); $i++) {{
      $out.ChangeLink($links.GetValue($i),$out.FullName,1)
    }}
  }}
  $out.Save()
  $g.Close($false); $old.Close($false); $out.Close($false)
}} finally {{ $x.Quit() }}
"""
    enc = base64.b64encode(ps.encode("utf-16-le")).decode()
    r = subprocess.run(["powershell", "-NoProfile", "-EncodedCommand", enc], capture_output=True, timeout=300)
    if r.returncode != 0:
        sys.exit("Excel merge failed:\n" + r.stderr.decode("utf-8", "replace")[-2000:])


def external_refs(wb):
    """Formulas still pointing at another workbook ('[...]'), as 'Sheet!A1'."""
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = getattr(c.value, "text", c.value)
                if isinstance(v, str) and v.startswith("=") and "[" in v:
                    bad.append(f"{ws.title}!{c.coordinate}")
    return bad


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--force", action="store_true", help="overwrite an existing MaxPain_GEX.xlsx")
    if OUT.exists() and not ap.parse_args().force:
        sys.exit(f"{OUT.name} already exists - use --force to rebuild it")
    excel_merge()
    wb = openpyxl.load_workbook(OUT)
    bad = external_refs(wb)
    if bad:
        sys.exit(f"merge left formulas pointing at other workbooks: {bad[:5]}")
    # Excel's cross-workbook copy leaves a link part to GEX V.4.1 that nothing references
    # (checked above); dropping it stops Excel asking to "update links" on every open.
    wb._external_links = []
    extend_layout(wb)
    rebuild_charts(wb)
    save_atomic(wb, OUT)
    print(f"Built {OUT.name}: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
