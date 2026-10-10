"""After each logged contract expires, record what the futures price actually did next to
what Max Pain / GEX said beforehand, in an 'Outcome' sheet of the Max Pain workbook.

CME only keeps ~5 trading days, so this has to run daily. A contract whose days have partly
fallen out of that window is still recorded, and the 'ข้อมูลครบ (วัน)' column shows how many
days were actually found. analyze.py reads the sheet.
"""
import json
import os
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from charts import rebuild_charts
from fetch_cme import MONTHS, WEEKDAY, _monthly_dates, cme_session, expiry_date, futures_row, px
from fetch_gamma import qs_code
from update_gex import AGG_LOG, GEX_LOG
from util import save_atomic

HERE = Path(__file__).parent
CONFIG = Path(os.environ.get("MAXPAIN_CONFIG", HERE / "config.json"))
TRADE_DATES = "/CmeWS/mvc/Volume/TradeDates?exchange=CBOT&isProtected"
OUT_HEAD = ["วันที่ข้อมูล", "สัญญา", "วันหมดอายุ", "ราคาเริ่ม", "Max Pain", "NET GEX", "Gamma Flip",
            "Settle วันหมดอายุ", "High ช่วงถือ", "Low ช่วงถือ", "|เริ่ม-MP|", "|settle-MP|",
            "|settle-เริ่ม|", "เข้าหา MP", "Range %", "ข้อมูลครบ (วัน)", "บันทึกเมื่อ",
            "Futures ที่ใช้เป็นราคา",
            # for paper trades (rules.py): data for day T only exists after T closes, so the
            # realistic entry is the next trading day's open, not T's settle
            "Entry (open วันถัดไป)", "Change วันเริ่ม", "โหมด GEX", "ความชัดเจน",
            "Call Wall", "Put Wall", "Pin #1",
            # the all-expiration GEX ('Log รวม') for the same data date, to compare against
            # the single-contract reading above (rules.R1_ALL / R2_ALL)
            "โหมด (รวม)", "ความชัดเจน (รวม)", "Call Wall (รวม)", "Put Wall (รวม)"]
OUT_FMT = ["yyyy-mm-dd", None, "yyyy-mm-dd", "0.0", "0", "0", "0", "0.0", "0.0", "0.0",
           "0.0", "0.0", "0.0", None, "0.00%", None, "yyyy-mm-dd hh:mm", None,
           "0.0", "+0.0;-0.0", None, "0.00", "0", "0", "0",
           None, "0.00", "0", "0"]
# Log rows from before the futures month was derived automatically have no month; their
# start price used the old OI>10,000 guess, so the outcome uses the same guess (a consistent
# price move) but distances to Max Pain can be off by a futures-month spread. analyze.py
# leaves these rows out.
GUESSED = "เดา: "


def parse_contract(s):
    """'MW1 Week 4 - SEP 2026 (U26)' -> ('MW1', 'Week 4 - SEP 2026');
    'AME OCT 2026 (V26)' -> ('AME', 'OCT 2026') - the standard monthly series; None otherwise."""
    m = re.match(r"(\w+) (Week \d - [A-Z]{3} \d{4}) \(", s or "")
    if m and m.group(1) in WEEKDAY:
        return m.group(1), m.group(2)
    m = re.match(r"AME ([A-Z]{3} \d{4}) \(", s or "")
    if m and m.group(1)[:3] in MONTHS:
        return "AME", m.group(1)
    return None


def monthly_key(label):
    """'OCT 2026' -> (10, 2026): the numeric (month, year) fetch_cme._monthly_dates is keyed by
    (the two CME endpoints disagree on the label text but agree on these numbers)."""
    return MONTHS.index(label[:3]) + 1, int(label[4:8])


def trading_days(start, end, available=None):
    """Weekdays d with start < d <= end. With `available` (CME's own trade dates), a weekday
    inside CME's window that has no trade date is an exchange holiday and is left out - so
    'ข้อมูลครบ' doesn't count a holiday as a missing day. Days outside the window are kept."""
    lo, hi = (min(available), max(available)) if available else (None, None)
    days, d = [], start + timedelta(days=1)
    while d <= end:
        holiday = available and lo <= d <= hi and d not in available
        if d.weekday() < 5 and not holiday:
            days.append(d)
        d += timedelta(days=1)
    return days


def outcome_metrics(start, mp, settle, hi, lo):
    d_start, d_end = abs(start - mp), abs(settle - mp)
    return {
        "d_start": d_start,
        "d_end": d_end,
        "move": abs(settle - start),
        "toward": d_end < d_start,
        "range_pct": (hi - lo) / start if hi is not None and lo is not None else None,
    }


def _as_date(v):
    return v.date() if isinstance(v, datetime) else v


def _gex_by_key(cfg):
    """{(trade date, QuikStrike code): GEX Log row, padded to 14 columns} - rows logged
    before the level columns existed just have None there."""
    path = HERE / cfg["gex_workbook"]
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if GEX_LOG not in wb.sheetnames:
            return {}
        return {(_as_date(r[0]), r[1]): tuple(r) + (None,) * (14 - len(r))
                for r in wb[GEX_LOG].iter_rows(min_row=2, values_only=True) if r[0]}
    finally:
        wb.close()


def _agg_by_date(cfg):
    """{trade date: (mode, conviction, call wall, put wall)} from 'Log รวม' (AGG_HEAD cols 6-9)."""
    path = HERE / cfg["gex_workbook"]
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if AGG_LOG not in wb.sheetnames:
            return {}
        return {_as_date(r[0]): tuple(r[6:10]) for r in wb[AGG_LOG].iter_rows(min_row=2, values_only=True)
                if r[0] and len(r) >= 10}
    finally:
        wb.close()


def main():
    cfg = json.load(open(CONFIG, encoding="utf-8"))
    path = HERE / cfg["workbook"]
    wb = openpyxl.load_workbook(path)
    if "Outcome" not in wb.sheetnames:
        wb.create_sheet("Outcome").append(OUT_HEAD)
    out = wb["Outcome"]
    for i, h in enumerate(OUT_HEAD, start=1):  # older, shorter header: extend in place
        if out.cell(1, i).value is None:
            out.cell(1, i).value = h
    done = {(_as_date(r[0]), r[1]) for r in out.iter_rows(min_row=2, values_only=True) if r[0]}

    pending, monthly = [], []
    for r in wb["Log"].iter_rows(min_row=2, values_only=True):
        parsed = parse_contract(r[1])
        if not parsed or r[3] is None or (_as_date(r[0]), r[1]) in done:
            continue
        if parsed[0] == "AME":
            # expiry only known from CME's quotes page (inside the session below), which lists
            # upcoming contracts only - so just the recent ones, or a row CME has already
            # dropped would open the browser every day for nothing
            if (date.today() - _as_date(r[0])).days <= 60:
                monthly.append((r, parsed))
            continue
        exp = expiry_date(*parsed)
        if exp < date.today():  # skips the browser entirely on days with nothing to record
            pending.append((r, parsed, exp))
    if not pending and not monthly:
        print("Outcomes: nothing new to record")
        return

    gex, agg = _gex_by_key(cfg), _agg_by_date(cfg)
    added = []
    with cme_session(cfg) as get:
        available = {datetime.strptime(d["tradeDate"], "%Y%m%d").date() for d in get(TRADE_DATES)}
        if monthly:
            ends = _monthly_dates(get, cfg)
            for r, parsed in monthly:
                exp = ends.get(monthly_key(parsed[1]))
                if exp and exp < date.today():
                    pending.append((r, parsed, exp))
        cache = {}

        def row_for(d, month):
            if d not in available:
                return None
            if (d, month) not in cache:
                cache[d, month] = futures_row(get, {**cfg, "price_month": month}, d.strftime("%Y%m%d"))
            return cache[d, month]

        for r, (family, label), exp in pending:
            if exp > max(available):
                continue  # CME hasn't published the expiry day yet - try again tomorrow
            log_date, month = _as_date(r[0]), r[11] if len(r) > 11 else None
            days = trading_days(log_date, exp, available)
            got = [g for g in (row_for(d, month) for d in days) if g]
            exp_row = row_for(exp, month)
            settle = px(exp_row["settle"]) if exp_row else None
            highs = [h for h in (px(g["high"]) for g in got) if h is not None]
            lows = [lo for lo in (px(g["low"]) for g in got) if lo is not None]
            hi, lo = (max(highs), min(lows)) if highs and lows else (None, None)
            start, mp = r[3], r[4]
            m = outcome_metrics(start, mp, settle, hi, lo) if settle is not None else {}
            # no QuikStrike code for the monthly series -> no single-contract GEX to join
            g = gex.get((log_date, qs_code(family, label)), (None,) * 14) if family != "AME" else (None,) * 14
            a = agg.get(log_date, (None,) * 4)
            used = month or (GUESSED + exp_row["month"] if exp_row else GUESSED.strip())
            first = row_for(days[0], month) if days else None
            entry = px(first["open"]) if first else None
            day0 = row_for(log_date, month)
            change = px(day0["change"]) if day0 else None
            out.append([log_date, r[1], exp, start, mp, g[5], g[7], settle, hi, lo,
                        m.get("d_start"), m.get("d_end"), m.get("move"), m.get("toward"),
                        m.get("range_pct"), f"{len(got)}/{len(days)}", datetime.now(), used,
                        entry, change, g[6], g[13], g[10], g[11], g[12], *a])
            for cell, fmt in zip(out[out.max_row], OUT_FMT):
                if fmt:
                    cell.number_format = fmt
            added.append(f"{r[1]} settle={settle} MP={mp} ({len(got)}/{len(days)} days)")

    if not added:
        print("Outcomes: nothing new to record (expiry days not published by CME yet)")
        return
    rebuild_charts(wb)  # openpyxl drops charts on save
    try:
        save_atomic(wb, path)
    except PermissionError:
        sys.exit(f"Cannot save - close {path.name} in Excel and run again")
    for a in added:
        print("Outcome:", a)


if __name__ == "__main__":
    main()
