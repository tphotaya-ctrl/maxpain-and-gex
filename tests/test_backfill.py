"""backfill.sort_by_date on an in-memory sheet - no network."""
from datetime import datetime

import openpyxl

from backfill import sort_by_date, target_for


def test_target_is_the_next_day_like_the_morning_run():
    assert target_for("20260929") == "2026-09-30"  # Tue data -> the Wed contract, not Tue's (expired, 0 OI)
    assert target_for("20261002") == "2026-10-03"  # Fri -> Sat; next-expiry fallback then gives Mon


def _sheet(dates):
    ws = openpyxl.Workbook().active
    ws.append(["วันที่", "สัญญา"])
    for d in dates:
        ws.append([d, f"c{d.day}"])
    return ws


def test_backfilled_older_day_moves_before_newer_ones():
    ws = _sheet([datetime(2026, 9, 25), datetime(2026, 10, 2), datetime(2026, 9, 29)])
    assert sort_by_date(ws) is True
    assert [r[1] for r in ws.iter_rows(min_row=2, values_only=True)] == ["c25", "c29", "c2"]
    assert ws.cell(1, 1).value == "วันที่"  # header untouched


def test_already_sorted_is_left_alone_and_same_day_keeps_order():
    ws = _sheet([datetime(2026, 9, 25), datetime(2026, 9, 25), datetime(2026, 9, 30)])
    ws.cell(3, 2).value = "second"
    assert sort_by_date(ws) is False
    assert ws.cell(3, 2).value == "second"
