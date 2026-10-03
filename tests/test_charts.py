"""charts.rebuild_charts on small in-memory workbooks - no files, no Excel."""
import openpyxl

from charts import filled, rebuild_charts


def _combined(oi_rows, gamma_rows):
    wb = openpyxl.Workbook()
    wb.active.title = "OI Data"
    for name in ("Max Pain Calc", "Gamma Data", "GEX Calc"):
        wb.create_sheet(name)
    for i in range(oi_rows):
        wb["OI Data"].cell(2 + i, 1).value = 4000 + 5 * i
    for i in range(gamma_rows):
        wb["Gamma Data"].cell(2 + i, 1).value = 4000 + 5 * i
    return wb


def test_filled_counts_rows_from_2():
    wb = _combined(3, 0)
    assert filled(wb["OI Data"], 201) == 3


def test_rebuild_charts_puts_back_both_and_sizes_them():
    wb = _combined(4, 2)
    rebuild_charts(wb)
    mp, gex = wb["Max Pain Calc"]._charts, wb["GEX Calc"]._charts
    assert len(mp) == 1 and len(gex) == 1
    assert gex[0].series[0].val.numRef.f == "'GEX Calc'!$C$2:$C$3"
    assert mp[0].series[0].val.numRef.f == "'Max Pain Calc'!$B$2:$B$5"


def test_rebuild_charts_is_idempotent_and_skips_empty_data():
    wb = _combined(4, 0)
    rebuild_charts(wb)
    rebuild_charts(wb)
    assert len(wb["Max Pain Calc"]._charts) == 1
    assert wb["GEX Calc"]._charts == []


def test_rebuild_charts_ignores_missing_sheets():
    wb = openpyxl.Workbook()
    rebuild_charts(wb)  # neither pair present: no error
