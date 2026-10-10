"""Pure-logic tests for update_gex.py - no network, no Excel.

Expected values were cross-checked by hand and by running the functions directly before
being written into these assertions.
"""
import pytest

from update_gex import _mode, call_wall, conviction, gamma_flip, gex_status, put_wall, select_rows, top_pin

# strike, call, put - price 112
LEVELS = [(100, 0, 10), (105, 2, 9), (110, 5, 5), (115, 12, 3), (120, 8, 1), (125, 12, 0)]


def test_call_wall_mirrors_j9():
    # >= 112 and call dominant: 115 (12), 120 (8), 125 (12) -> largest 12, lowest strike wins the tie
    assert call_wall(LEVELS, 112) == 115


def test_put_wall_mirrors_j10():
    # <= 112 and put dominant: 100 (10), 105 (9); 110 is 5 vs 5, not 1.5x -> 100
    assert put_wall(LEVELS, 112) == 100


def test_walls_none_without_price_or_candidates():
    assert call_wall(LEVELS, None) is None
    assert put_wall([(100, 5, 5)], 100) is None


def test_conviction_mirrors_j18():
    # net -10-7+0+9+7+12 = 11; sum|c-p| = 45
    assert conviction(LEVELS) == pytest.approx(11 / 45)
    assert conviction([(100, 3, 3)]) is None


def test_top_pin_mirrors_j30_including_row_tiebreak():
    # unit = 112*0.005 = 0.56; 110: 10/(2/0.56) = 2.8, 115: 15/(3/0.56) = 2.8 -> tie, the later
    # row (115) wins via the sheet's ROW()/1e9 term; 100 and 125 are outside 10% window
    assert top_pin(LEVELS, 112, window=0.1, min_dist=0) == 115
    # sheet defaults (5% window, 25 min distance): nothing qualifies on this tiny table
    assert top_pin(LEVELS, 112) is None
    assert top_pin(LEVELS, None) is None


def test_gex_status_follows_v41_j5_rule():
    # net +8 of spread 12 (67%) -> clear positive; mirrored -> negative
    assert gex_status([(100, 10, 0), (105, 0, 2)]) == "Positive GEX (นิ่ง)"
    assert gex_status([(100, 0, 10), (105, 2, 0)]) == "Negative GEX (วิ่ง)"
    # net 0 of spread 20 -> under 5% -> no mode; all-zero table -> no mode too
    assert gex_status([(100, 10, 0), (105, 0, 10)]) == "ไม่มีโหมด"
    assert gex_status([(100, 0, 0)]) == "ไม่มีโหมด"


def test_gex_status_small_net_is_no_mode():
    # NET +2 against a spread of 100 is 2% -> below the 5% threshold
    assert gex_status([(100, 50, 0), (105, 0, 48), (110, 1, 1)]) == "ไม่มีโหมด"


def test_mode_reads_old_and_new_log_wording():
    assert _mode("Positive GEX (นิ่ง)") == "+"
    assert _mode("Negative GEX (แกว่งแรง)") == "-"  # pre-V.4.1 Log rows
    assert _mode("Negative GEX (วิ่ง)") == "-"
    assert _mode("ไม่มีโหมด") is None
    assert _mode(None) is None


def test_select_rows_includes_one_zero_strike_each_side():
    g = {100: (0, 0), 105: (5, 2), 110: (2, 8), 115: (0, 0)}
    assert select_rows(g) == [(100, 0, 0), (105, 5, 2), (110, 2, 8), (115, 0, 0)]


def test_select_rows_all_zero_raises():
    with pytest.raises(RuntimeError):
        select_rows({100: (0, 0), 105: (0, 0)})


def test_gamma_flip_finds_sign_change():
    rows = select_rows({100: (0, 0), 105: (5, 2), 110: (2, 8), 115: (0, 0)})
    # cumulative net gamma: 0 -> +3 (at 105) -> -3 (at 110, sign change) -> stays -3 (at 115)
    assert gamma_flip(rows, price=108) == 110
    assert gamma_flip(rows, price=100) == 110  # only one flip point either way


def test_gamma_flip_none_when_cumulative_never_changes_sign():
    rows = select_rows({100: (5, 1), 105: (3, 1)})  # always net-positive
    assert gamma_flip(rows, price=102) is None


def test_ref_price_prefers_live_unless_settle_configured():
    from update_gex import ref_price
    d = {"price": 4187.1, "live_price": 4168.4}
    assert ref_price({}, d) == 4168.4
    assert ref_price({"gex_price_source": "settle"}, d) == 4187.1
    assert ref_price({}, {"price": 4187.1, "live_price": None}) == 4187.1


def test_aggregate_sums_per_strike_across_expirations():
    from update_gex import aggregate
    per_code = {"OG2V6": {4100: (1, 5), 4200: (3, 0)}, "OGX6": {4200: (10, 2), 4300: (4, 0)}}
    assert aggregate(per_code) == [(4100, 1.0, 5.0), (4200, 13.0, 2.0), (4300, 4.0, 0.0)]


def test_agg_levels_apply_the_workbook_rules():
    from update_gex import agg_levels as levels
    rows = [(4000, 0, 40), (4050, 5, 30), (4100, 20, 18), (4150, 30, 2), (4200, 60, 10), (4250, 25, 0)]
    lv = levels(rows, 4120)
    assert lv["call_wall"] == 4200          # biggest call >= 1.5x put at/above price
    assert lv["put_wall"] == 4000           # biggest put >= 1.5x call at/below price; 4100 is mixed
    assert lv["peak"] == 4200 and lv["net"] == 140 - 100
    assert lv["flip"] == 4200               # cumulative -40,-65,-63,-35,+15: turns positive at 4200
    assert lv["mode"].startswith("Positive")


def _gex_clone(tmp_path, monkeypatch):
    import shutil
    from pathlib import Path

    import update_gex
    src = Path(__file__).resolve().parent.parent / "MaxPain_GEX.xlsx"
    dst = tmp_path / "wb.xlsx"
    shutil.copy(src, dst)
    monkeypatch.setattr(update_gex, "notify", lambda *a, **k: None)  # no popup/Telegram from tests
    return {"gex_workbook": str(dst), "gex_price_source": "settle"}, dst


def test_update_gex_uses_the_matrix_table_without_touching_quikstrike(tmp_path, monkeypatch, capsys):
    from datetime import date

    import openpyxl

    import update_gex
    cfg, path = _gex_clone(tmp_path, monkeypatch)

    def no_popup_path(*a, **k):
        raise AssertionError("fetch_gamma must not run when the matrix already has the contract")
    monkeypatch.setattr(update_gex, "fetch_gamma", no_popup_path)
    g = {4100: (0, 0), 4150: (5, 20), 4200: (30, 4), 4250: (0, 0)}
    d = {"family": "MW1", "label": "Week 2 - OCT 2026", "trade_date": date(2026, 10, 8), "price": 4157.0}
    out = update_gex.update_gex(cfg, d, g=g)
    assert (out["calls"], out["puts"], out["code"]) == (35, 24, "G2MV6")
    assert "source=matrix" in capsys.readouterr().out
    log = openpyxl.load_workbook(path)[update_gex.GEX_LOG]
    last = [c.value for c in log[log.max_row]]
    assert last[1] == "G2MV6" and last[2] == 4157.0 and last[5] == 11


def test_update_gex_all_writes_log_ruam_and_replaces_the_same_day(tmp_path, monkeypatch):
    from datetime import date

    import openpyxl

    import fetch_gamma
    import update_gex
    cfg, path = _gex_clone(tmp_path, monkeypatch)
    per_code = {"G2MV6": {4150: (5, 20), 4200: (30, 4)}, "OGX6": {4200: (10, 2), 4300: (4, 0)}}
    monkeypatch.setattr(fetch_gamma, "fetch_gamma_all", lambda cfg, td: per_code)
    d = {"family": "MW1", "label": "Week 2 - OCT 2026", "trade_date": date(2026, 10, 8), "price": 4157.0}
    out = update_gex.update_gex_all(cfg, d)
    assert out["per_code"] is per_code and out["net"] == (5 + 40 + 4) - (20 + 6)
    update_gex.update_gex_all(cfg, d)  # same trade date again
    rows = [r for r in openpyxl.load_workbook(path)[update_gex.AGG_LOG].iter_rows(min_row=2, values_only=True)
            if r[0] and r[0].date() == date(2026, 10, 8)]
    assert len(rows) == 1 and rows[0][1] == "G2MV6, OGX6"
