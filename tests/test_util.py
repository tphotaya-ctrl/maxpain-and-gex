"""Tests for util.save_atomic - uses real files under pytest's tmp_path, no network."""
import os

import openpyxl
import pytest

from util import save_atomic


def _wb(value):
    wb = openpyxl.Workbook()
    wb.active["A1"] = value
    return wb


def test_save_atomic_writes_and_cleans_up_temp_file(tmp_path):
    path = tmp_path / "book.xlsx"
    save_atomic(_wb("hello"), path)
    assert path.exists()
    assert openpyxl.load_workbook(path).active["A1"].value == "hello"
    assert not (tmp_path / ".tmp_book.xlsx").exists()


def test_save_atomic_overwrites_existing_file(tmp_path):
    path = tmp_path / "book.xlsx"
    save_atomic(_wb("first"), path)
    save_atomic(_wb("second"), path)
    assert openpyxl.load_workbook(path).active["A1"].value == "second"


@pytest.mark.skipif(os.name != "nt", reason="relies on Windows' exclusive-lock file semantics")
def test_save_atomic_raises_and_leaves_original_intact_when_target_is_locked(tmp_path):
    path = tmp_path / "book.xlsx"
    save_atomic(_wb("original"), path)
    handle = open(path, "rb")  # holding this open blocks os.replace() on Windows
    try:
        with pytest.raises(PermissionError):
            save_atomic(_wb("blocked"), path)
    finally:
        handle.close()
    assert openpyxl.load_workbook(path).active["A1"].value == "original"
    assert not (tmp_path / ".tmp_book.xlsx").exists()  # temp file cleaned up even on failure
