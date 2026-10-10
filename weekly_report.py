"""Once a week, a short Telegram message with the running evidence: sample size vs the 30
needed, Max Pain vs the naive "price stays put" guess, each paper rule vs the baseline, and
how many days of the week actually got GEX. Same numbers as analyze.py, no need to open it.

Sent by the daily run from Friday on (or, if Friday's run never happened, by the next
week's first run, for the week before); weekly_sent.txt (git-ignored) keeps it to once.

    python weekly_report.py          # print it (and send if due)
    python weekly_report.py --force  # send now regardless
"""
import json
import os
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

from analyze import MIN_N, load_records, paper_trades, summarize
from update_gex import GEX_LOG

HERE = Path(__file__).parent
CONFIG = Path(os.environ.get("MAXPAIN_CONFIG", HERE / "config.json"))
STAMP = HERE / "weekly_sent.txt"


def target_week(today):
    """The ISO week a report is due for: this week from Friday on, otherwise last week's."""
    d = today if today.weekday() >= 4 else today - timedelta(days=7)
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def week_bounds(week):
    """'2026-W41' -> (Monday, Friday) of that ISO week."""
    y, w = week.split("-W")
    monday = date.fromisocalendar(int(y), int(w), 1)
    return monday, monday + timedelta(days=4)


def due(today, stamp=None):
    try:
        return (stamp or STAMP).read_text().strip() != target_week(today)
    except OSError:
        return True


def gex_coverage(log_dates, gex_dates, week):
    """(days with data, of which with GEX) for trade dates inside `week`."""
    mon, fri = week_bounds(week)
    days = {d for d in log_dates if mon <= d <= fri}
    return len(days), len(days & set(gex_dates))


def format_weekly(week, records, mid, edge, coverage):
    s = summarize(records, mid, edge)
    lines = [f"สรุปรายสัปดาห์ {week}"]
    n = s["n"]
    lines.append(f"ตัวอย่างผลจริง: {n}/{MIN_N} สัญญา" + (" (ยังเป็น noise)" if n < MIN_N else ""))
    if n:
        better = "Max Pain แม่นกว่า" if s["mae_mp"] < s["mae_naive"] else "เดาว่าราคาไม่ขยับแม่นกว่า"
        lines.append(f"Max Pain คลาด {s['mae_mp']:.1f} จุด vs เดาไม่ขยับ {s['mae_naive']:.1f} -> {better}")
        lines.append(f"ราคาเข้าหา Max Pain {s['toward_rate']:.0%} ของสัญญา")
    lines.append("")
    lines.append("เทรดจำลองสะสม:")
    for name, st in paper_trades(records).items():
        if st["n"]:
            lines.append(f"- {name}: {st['n']} เทรด ชนะ {st['win_rate']:.0%} รวม {st['total']:+.1f} จุด")
        else:
            lines.append(f"- {name}: ยังไม่มีเทรด")
    days, with_gex = coverage
    lines.append("")
    lines.append(f"GEX สัปดาห์นี้: ได้ {with_gex}/{days} วัน" if days else "สัปดาห์นี้ไม่มีข้อมูลบันทึก")
    return "\n".join(lines)


def _dates(ws):
    return [r[0].date() if isinstance(r[0], datetime) else r[0]
            for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def build(cfg, week):
    loaded = load_records(cfg)
    records, _, mid, edge = loaded if loaded else ([], 0, 0.02, 0.05)
    wb = openpyxl.load_workbook(HERE / cfg["workbook"], read_only=True)
    try:
        log_dates = _dates(wb["Log"]) if "Log" in wb.sheetnames else []
        gex_dates = _dates(wb[GEX_LOG]) if GEX_LOG in wb.sheetnames else []
    finally:
        wb.close()
    return format_weekly(week, records, mid, edge, gex_coverage(log_dates, gex_dates, week))


def main(force=False, today=None):
    today = today or date.today()
    if not (force or due(today)):
        return None
    week = target_week(today)
    text = build(json.load(open(CONFIG, encoding="utf-8")), week)
    print(text)
    from telegram_report import send_text
    if send_text(text):  # stamp only once it's actually delivered - no Telegram, no stamp
        STAMP.write_text(week)
    return text


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main(force="--force" in sys.argv)
